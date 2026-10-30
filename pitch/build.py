"""Builds pitch/gatekeeper-pitch.pptx: 7 slides mirroring video/narration.txt, notes verbatim from it.
Run: python pitch/build.py   (needs python-pptx 1.0.2)"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
STILLS = HERE / "stills"
UI, SERIF, MONO = "Atkinson Hyperlegible Next", "Source Serif 4", "IBM Plex Mono"
BG, CARD, BORDER, INK, MUTED = "EDF0F2", "FFFFFF", "C9D0D6", "16202A", "5B6670"
BLUE, RED, GREEN = "1D4E89", "9B2226", "2D6A4F"

narration = (HERE.parent / "video" / "narration.txt").read_text(encoding="utf-8")
SEG = {int(m.group(1)): m.group(2).strip() for m in re.finditer(r"^\[(\d)\] [^\n]*\n(.+?)(?=\n\n|\Z)", narration, re.M | re.S)}
assert len(SEG) == 6, SEG.keys()

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def rect(slide, x, y, w, h, fill=CARD, line=BORDER, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, runs, size=18, font=UI, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, overrides) tuples."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for chunk in (para if isinstance(para, list) else [(para, {})]):
            t, o = chunk if isinstance(chunk, tuple) else (chunk, {})
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold = o.get("font", font), Pt(o.get("size", size)), o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def wordmark(slide, x, y, size):
    return text(slide, x, y, 8, size / 50, [[(".", {"color": BLUE, "bold": True}), ("Gatekeeper", {"bold": True})]], size=size, font=UI, anchor=MSO_ANCHOR.TOP)


def base(title, note):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 7.5, fill=BG, line=None)
    if title:
        text(s, 0.6, 0.45, 12.1, 0.9, title, size=34, font=SERIF, bold=True)
        wordmark(s, 11.4, 6.85, 14)
    s.notes_slide.notes_text_frame.text = note
    return s


def chip(slide, x, y, w, label, color):
    c = rect(slide, x, y, w, 0.5, fill=color, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = MONO, Pt(14), True
    r.font.color.rgb = rgb("FFFFFF")
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def line(slide, x1, y1, x2, y2, color=BLUE, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)


def picture(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(STILLS / name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


# 1 Title
s = base(None, "No narration on the title slide. The spoken pitch starts on slide 2.")
wordmark(s, 2.0, 2.4, 72)
text(s, 2.0, 3.9, 10, 0.8, "Decision support for the European Health Data Space", size=26, font=SERIF, color=MUTED)
text(s, 2.0, 5.9, 10, 0.5, "AIM Healthcare Hackathon Vienna 2026, Track 3. Oguz Nacar and Arda Sarikaya.", size=14, color=MUTED)

# 2 Problem
s = base("Reliable health data is expensive and complicated to access", SEG[1])
rect(s, 0.6, 1.8, 5.6, 4.6)
text(s, 1.0, 2.2, 4.8, 1.6, "EUR 4,000", size=60, font=MONO, color=RED, bold=True)
text(s, 1.0, 3.9, 4.8, 2.2, ["Standard processing fee for one application in Germany.", "Paid before anyone knows whether the request can be granted."], size=18, color=INK, spacing=10)
text(s, 6.8, 1.9, 6.0, 4.6, [
    [("Medical research depends on reliable health data.", {"bold": True})],
    "Every application goes through a Health Data Access Body under Regulation (EU) 2025/327.",
    "Applicants write the request, wait, and find out late which article blocks it.",
    [("That is why we created Gatekeeper.", {"bold": True, "color": BLUE})],
], size=20, spacing=14)

# 3 Applicant interface
s = base("Applicant interface: submit, then see the decision", SEG[2])
picture(s, "applicant-form.png", 0.6, 1.6, w=7.2)
picture(s, "notice-approve-spe.png", 8.6, 1.6, h=4.9)
text(s, 0.6, 6.2, 7.2, 0.5, "Request form: purpose, question, definitions.", size=13, color=MUTED)
text(s, 8.6, 6.6, 4.2, 0.5, "Decision notice: APPROVE_SPE with Art. 53, 66(3), 68, 73(2).", size=12, color=MUTED)

# 4 AI decision engine
s = base("AI decision engine: structured, consistent, explainable", SEG[3])
steps = [("1", "Each request is analyzed", "Purpose, question and definitions are read as structured fields."),
         ("2", "Relevant requirements are checked", "Chapter IV of Regulation (EU) 2025/327, rule by rule."),
         ("3", "An explainable recommendation", "Every outcome names the rule applied and cites the articles.")]
for i, (num, head, body) in enumerate(steps):
    x = 0.6 + i * 4.15
    rect(s, x, 1.7, 3.85, 2.9)
    c = rect(s, x + 0.3, 2.0, 0.6, 0.6, fill=BLUE, line=None, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = num
    r.font.name, r.font.size, r.font.bold = MONO, Pt(16), True
    r.font.color.rgb = rgb("FFFFFF")
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(s, x + 0.3, 2.8, 3.3, 0.7, head, size=18, bold=True)
    text(s, x + 0.3, 3.5, 3.3, 1.0, body, size=14, color=MUTED)
text(s, 0.6, 5.0, 12, 0.5, "Three possible outcomes", size=16, bold=True)
chip(s, 0.6, 5.55, 2.4, "REJECT", RED)
chip(s, 3.2, 5.55, 3.6, "APPROVE_STATISTICAL", BLUE)
chip(s, 7.0, 5.55, 2.6, "APPROVE_SPE", GREEN)
text(s, 0.6, 6.25, 12, 0.5, "Specialized decision models instead of a free-form LLM answer: the same request gets the same decision.", size=14, color=MUTED)

# 5 Staff interface and federated data
s = base("Staff interface and federated data", SEG[4])
picture(s, "staff-casefile.png", 0.6, 1.6, w=6.0)
text(s, 0.6, 5.75, 6.0, 0.4, "Staff case file: rule applied, articles, who read the request, numbers.", size=12, color=MUTED)
dx, dy = 7.2, 1.7
for i, name in enumerate(["Klinikum Nord", "Klinikum Süd", "Klinikum West"]):
    y = dy + i * 1.15
    rect(s, dx, y, 2.3, 0.8)
    text(s, dx + 0.15, y, 2.0, 0.8, name, size=13, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    line(s, dx + 2.3, y + 0.4, dx + 3.5, dy + 1.55 + 0.4)
text(s, dx + 3.5, dy + 2.45, 2.4, 0.4, "totals only, no patient rows", size=11, color=BLUE, align=PP_ALIGN.CENTER)
rect(s, dx + 3.5, dy + 1.55, 2.4, 0.8, fill=BLUE, line=None)
text(s, dx + 3.5, dy + 1.55, 2.4, 0.8, "Health Data Access Body", size=13, bold=True, color="FFFFFF", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, dx, 5.2, 5.9, 1.2, [
    [("Queries run inside each hospital. ", {"bold": True}), "Patient rows never leave; only totals are merged centrally."],
    [("Opted-out patients excluded. ", {"bold": True}), "Counts under 10 shown as <10."],
], size=14, spacing=8)

# 6 Results
s = base("Results on the 10 sample requests", SEG[5])
tiles = [("10/10", "correct decisions", BLUE), ("6/6", "correct numbers", BLUE), ("0", "privacy leaks", GREEN),
         ("10/10", "consistent over 3 runs", BLUE), ("12.7 s", "slowest request", INK)]
for i, (num, lab, col) in enumerate(tiles):
    x = 0.6 + i * 2.47
    rect(s, x, 2.0, 2.27, 2.8)
    text(s, x + 0.1, 2.35, 2.07, 1.2, num, size=36 if len(num) > 3 else 48, font=MONO, color=col, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 3.7, 1.9, 0.8, lab, size=14, color=MUTED, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.2, 12, 0.5, "Measured on the 10 sample requests, synthetic data.", size=14, color=MUTED)
text(s, 0.6, 5.7, 12, 0.8, "Decision accuracy, statistical correctness, privacy protection, consistency, processing time.", size=16)

# 7 Vision and limitations
s = base("Vision and limitations", SEG[6])
rect(s, 0.6, 1.7, 5.9, 3.6)
text(s, 1.0, 2.0, 5.1, 0.5, "Today", size=16, bold=True, color=MUTED)
text(s, 1.0, 2.6, 5.1, 2.5, ["Prototype runs the full workflow on synthetic hospital data.", "Three hospital databases in one process; in production each sits behind its own service.", "Cases live in memory; no login; runs locally."], size=15, spacing=8)
rect(s, 6.8, 1.7, 5.9, 3.6)
text(s, 7.2, 2.0, 5.1, 0.5, "Future work", size=16, bold=True, color=MUTED)
text(s, 7.2, 2.6, 5.1, 2.5, ["Integration with real European healthcare infrastructure.", "Help research institutions access and analyze health data more reliably, affordably, and securely."], size=15, spacing=8)
text(s, 0.6, 5.7, 12.1, 1.0, "Making European health research safer, smarter, and more accessible.", size=26, font=SERIF, bold=True, color=BLUE)

out = HERE / "gatekeeper-pitch.pptx"
prs.save(out)

EM = chr(0x2014)
# Self-check: 7 slides, every slide has notes, no em-dash anywhere in the deck text.
chk = Presentation(out)
assert len(chk.slides) == 7
for sl in chk.slides:
    assert sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip()
    for sh in sl.shapes:
        if sh.has_text_frame:
            assert EM not in sh.text_frame.text, sh.text_frame.text
    assert EM not in sl.notes_slide.notes_text_frame.text
print("ok", out)
